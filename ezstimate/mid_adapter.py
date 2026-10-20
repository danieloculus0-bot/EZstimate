"""Read historical MID forecast workbooks without interpreting stock as price.

Preserves source workbook/sheet/row, with no edits, no automatic costing,
no implicit import into the shop ERP.
"""
from __future__ import annotations

from datetime import date, datetime
from io import BytesIO
from pathlib import Path

MAX_ROWS = 20000
MAX_FILE = 16_000_000


def _text(raw):
    return str(raw).strip() if raw is not None else ""


def _date(raw):
    if isinstance(raw, (date, datetime)):
        return raw.date().isoformat() if isinstance(raw, datetime) else raw.isoformat()
    return _text(raw)


def _headings(row):
    return {_text(value).strip().lower(): column for column, value in enumerate(row) if _text(value)}


def _field(values, headings, *names):
    for name in names:
        pos = headings.get(name.lower())
        if pos is not None and pos < len(values):
            return values[pos]
    return None


def preview_mid(filename, content):
    """Recognize MID's demand event and buyer queue sheets (read-only)."""
    if Path(filename).suffix.lower() != ".xlsx" or not 0 < len(content) <= MAX_FILE:
        raise ValueError("Upload an MID XLSX workbook up to 16 MB")
    try:
        from openpyxl import load_workbook
        workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
    except Exception as exc:
        raise ValueError("Could not open XLSX workbook") from exc
    events, actions, warnings = [], [], []
    seen_formats = []
    try:
        for sheet in workbook:
            headings, kind = None, None
            for line_no, row in enumerate(sheet.iter_rows(values_only=True), 1):
                if line_no > MAX_ROWS:
                    warnings.append(f"{sheet.title}: only first {MAX_ROWS} rows considered")
                    break
                values = list(row)
                if line_no <= 30:
                    found = _headings(values)
                    if ("material part number" in found and
                        ("event type" in found or "needed source units" in found)):
                        headings, kind = found, "jobboss_demand_events"
                        if kind not in seen_formats:
                            seen_formats.append(kind)
                        continue
                    if ("part number" in found and "shortage qty" in found and
                        "risk level" in found):
                        headings, kind = found, "material_risk_register"
                        if kind not in seen_formats:
                            seen_formats.append(kind)
                        continue
                    if {"action", "material", "order by", "buy qty"}.issubset(found):
                        headings, kind = found, "buyer_action_queue"
                        if kind not in seen_formats:
                            seen_formats.append(kind)
                        continue
                if not headings or not any(_text(v) for v in values):
                    continue
                if kind == "jobboss_demand_events":
                    material = _text(_field(values, headings, "Normalized Key",
                                            "Material Part Number")).upper()
                    event = _text(_field(values, headings, "Event Type"))
                    if not material or not event:
                        continue
                    events.append({
                        "material_key": material, "event_type": event,
                        "need_window": _text(_field(values, headings, "Need Window")),
                        "need_start": _date(_field(values, headings, "Window Start")),
                        "need_end": _date(_field(values, headings, "Window End")),
                        "needed_source_units": _text(_field(values, headings, "Needed Source Units")),
                        "running_balance_source_units": _text(_field(
                            values, headings, "Running Balance Source Units")),
                        "running_balance_sheets": _text(_field(values, headings, "Running Balance (Sheets)")),
                        "risk_flag": _text(_field(values, headings, "Risk Flag")),
                        "job_number": _text(_field(values, headings, "Job Number")),
                        "po_number": _text(_field(values, headings, "PO Number")),
                        "snapshot_date": _date(_field(values, headings, "Snapshot Date")),
                        "source_file": Path(filename).name, "source_sheet": sheet.title,
                        "source_row": line_no,
                    })
                elif kind == "material_risk_register":
                    material = _text(_field(values, headings, "Part Number")).upper()
                    if not material:
                        continue
                    events.append({
                        "material_key": material, "event_type": "risk_register",
                        "shortage_quantity": _text(_field(values, headings, "Shortage Qty")),
                        "risk_flag": _text(_field(values, headings, "Risk Level")),
                        "need_start": _date(_field(values, headings, "Need Start")),
                        "need_end": _date(_field(values, headings, "Need End")),
                        "supplier": _text(_field(values, headings, "Supplier")),
                        "snapshot_date": _date(_field(values, headings, "Snapshot Date")),
                        "source_file": Path(filename).name, "source_sheet": sheet.title,
                        "source_row": line_no,
                    })
                elif kind == "buyer_action_queue":
                    material = _text(_field(values, headings, "Material")).upper()
                    action = _text(_field(values, headings, "Action"))
                    if not material or not action:
                        continue
                    actions.append({
                        "material_key": material, "action": action,
                        "order_by": _date(_field(values, headings, "Order By")),
                        "job_due": _date(_field(values, headings, "Job Due")),
                        "buy_qty": _text(_field(values, headings, "Buy Qty")),
                        "supplier": _text(_field(values, headings, "Supplier")),
                        "job": _text(_field(values, headings, "Controlling Job")),
                        "notes": _text(_field(values, headings, "Why Now"))[:300],
                        "source_file": Path(filename).name, "source_sheet": sheet.title,
                        "source_row": line_no,
                    })
    finally:
        workbook.close()
    if not events and not actions:
        raise ValueError("No recognized MID forecast, risk register, or buyer queue columns found")
    return {
        "schema": "ezstimate.mid_preview.v1",
        "source_file": Path(filename).name,
        "recognized_formats": seen_formats,
        "demand_event_count": len(events),
        "buyer_action_count": len(actions),
        "flagged_event_count": sum(bool(x.get("risk_flag")) for x in events),
        "buy_now_count": sum("BUY NOW" in x["action"].upper() for x in actions),
        "events": events[:150],
        "actions": actions[:150],
        "preview_truncated": len(events) > 150 or len(actions) > 150,
        "warnings": warnings,
        "quote_material_price": None,
        "persistent_import": False,
        "notes": ["Read-only preview, not a material purchase commitment",
                  "Job-demand events can repeat across windows; do not sum blindly",
                  "MID unit cost is not assumed to be a verified supplier invoice"],
    }
