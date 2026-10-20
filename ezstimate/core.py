"""Source-aware document candidate extraction and transparent quote arithmetic."""
from __future__ import annotations

import csv
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP
import hashlib
from io import BytesIO, StringIO
import re

SCHEMA = "ezstimate.quote_proposal.v1"
MONEY = Decimal("0.01")
FIELDS = ("material_unit_cost", "labor_hours_per_unit", "labor_rate_per_hour",
          "setup_hours", "setup_rate_per_hour", "outside_unit_cost",
          "scrap_rate", "overhead_rate", "margin_rate")
RATE_FIELDS = {"scrap_rate", "overhead_rate", "margin_rate"}
PATTERNS = {
    "material": re.compile(r"\b(?:MATERIAL|MATL|MAT\.?)\s*[:=-]\s*(.{3,60})", re.I),
    "drawing_number": re.compile(r"\b(?:DWG|DRAWING|PART)\s*(?:NO\.?|NUMBER|#)\s*[:=#-]?\s*([A-Z0-9][A-Z0-9._/-]{2,})", re.I),
    "purchase_order": re.compile(r"\b(?:P\.?O\.?|PURCHASE\s+ORDER)\s*(?:NO\.?|NUMBER|#)?\s*[:=#-]\s*([A-Z0-9][A-Z0-9._/-]{2,})", re.I),
    "quantity": re.compile(r"\b(?:QTY|QUANTITY)\s*[:=#-]\s*(\d+)\b", re.I),
}

def sha256(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()

def analyze_pdf(data: bytes, *, filename="document.pdf") -> dict:
    """Extract *candidate* fields from real selectable PDF text with source lines."""
    if not isinstance(data, bytes) or not data.startswith(b"%PDF-"):
        raise ValueError("PDF bytes are required")
    if len(data) > 20_000_000:
        raise ValueError("20MB PDF limit")
    import fitz
    candidates = []
    text_pages = 0
    with fitz.open(stream=data, filetype="pdf") as document:
        if not document.page_count:
            raise ValueError("PDF has no pages")
        for index, page in enumerate(document):
            txt = page.get_text("text")
            if txt.strip():
                text_pages += 1
            for line_no, line in enumerate(txt.splitlines(), 1):
                for kind, pattern in PATTERNS.items():
                    found = pattern.search(line)
                    if found:
                        value = found.group(1).strip().rstrip(".;,")
                        if value:
                            candidates.append({"kind": kind, "value": value,
                                "page": index + 1, "line": line_no, "raw_text": line.strip()[:200],
                                "review_required": True})
    return {"filename": filename, "sha256": sha256(data), "byte_size": len(data),
            "status": "review_required",
            "candidates": candidates,
            "warnings": ["No selectable PDF text; OCR required, not performed"] if text_pages == 0
                        else ([] if candidates else ["No recognizable title-block or PO field candidates"])}

def import_bom(data: bytes, *, filename: str) -> dict:
    """Read CSV/XLSX BOM rows with original line numbers; no fabricated costing."""
    if not isinstance(data, bytes) or len(data) > 10_000_000:
        raise ValueError("BOM bytes required (max 10MB)")
    ext = filename.lower().rsplit(".", 1)[-1]
    if ext == "csv":
        with StringIO(data.decode("utf-8-sig")) as source:
            raw = list(csv.reader(source))
    elif ext == "xlsx":
        from openpyxl import load_workbook
        book = load_workbook(BytesIO(data), data_only=True, read_only=True)
        try:
            raw = list(book.active.iter_rows(values_only=True))
        finally:
            book.close()
    else:
        raise ValueError("only CSV and XLSX BOMs are supported")
    if not raw or not raw[0]:
        raise ValueError("BOM header missing")
    headers = [str(h).strip() if h is not None else "" for h in raw[0]]
    if len(set(x.casefold() for x in headers)) != len(headers):
        raise ValueError("duplicate BOM headings")
    rows = []
    for row_number, values in enumerate(raw[1:], 2):
        if not any(x is not None and str(x).strip() for x in values):
            continue
        row = {h: ("" if val is None else str(val).strip())
               for h, val in zip(headers, values) if h}
        rows.append({"row": row_number, "cells": row, "review_required": True})
    return {"filename": filename, "sha256": sha256(data), "row_count": len(rows),
            "headers": headers, "rows": rows, "status": "review_required"}

def decimal_number(value, name: str) -> Decimal:
    if isinstance(value, bool) or value is None or not isinstance(value, (str,int,float,Decimal)):
        raise ValueError(f"explicit numeric input required: {name}")
    try:
        n = Decimal(str(value))
    except InvalidOperation as exc:
        raise ValueError(f"invalid numeric input: {name}") from exc
    if not n.is_finite() or n < 0:
        raise ValueError(f"{name} must be finite and nonnegative")
    return n

def calculate_quote(inputs: dict) -> dict:
    """Deterministic dollar estimate from explicit prices only, not inferred rates."""
    if not isinstance(inputs, dict) or "quantity" not in inputs:
        raise ValueError("quantity is required")
    quantity = decimal_number(inputs["quantity"], "quantity")
    if quantity != quantity.to_integral_value() or quantity < 1:
        raise ValueError("quantity must be a positive whole number")
    missing = sorted(set(FIELDS) - set(inputs))
    if missing:
        raise ValueError("missing explicit costs: " + ", ".join(missing))
    n = {field: decimal_number(inputs[field], field) for field in FIELDS}
    for field in RATE_FIELDS:
        if n[field] >= 1:
            raise ValueError(f"{field} must be below 1 (fractional rate)")
    material = n["material_unit_cost"] / (1 - n["scrap_rate"])
    machining = n["labor_hours_per_unit"] * n["labor_rate_per_hour"]
    setup = n["setup_hours"] * n["setup_rate_per_hour"] / quantity
    outside = n["outside_unit_cost"]
    direct = material + machining + setup + outside
    unit_cost = direct * (1 + n["overhead_rate"])
    sell = unit_cost / (1 - n["margin_rate"])
    def usd(val: Decimal) -> str:
        return str(val.quantize(MONEY, rounding=ROUND_HALF_UP))
    return {"quantity": int(quantity), "currency": "USD", "material_per_unit": usd(material),
            "machining_per_unit": usd(machining), "setup_per_unit": usd(setup),
            "outside_per_unit": usd(outside), "loaded_unit_cost": usd(unit_cost),
            "quoted_unit_price": usd(sell), "quoted_total": usd(sell * quantity),
            "formula": "material/(1-scrap) + hours*rate + setup*rate/qty + outside; apply overhead then divide by (1-margin)"}

def quote_proposal(quote_number: str, inputs: dict, *, drawing: dict | None = None,
                   bom: dict | None = None, po: dict | None = None,
                   released_drawing: dict | None = None) -> dict:
    if not isinstance(quote_number, str) or not quote_number.strip() or len(quote_number) > 100:
        raise ValueError("quote number required")
    if released_drawing is not None:
        if not drawing or released_drawing.get("sha256") != drawing.get("sha256"):
            raise ValueError("controlled release SHA does not match analyzed PDF")
        if released_drawing.get("lifecycle") != "Released" or not released_drawing.get("file_version_id"):
            raise ValueError("released controlled file version required")
    price = calculate_quote(inputs)
    warnings = [w for source in (drawing, po) if source for w in source.get("warnings", [])]
    return {"schema": SCHEMA, "quote_number": quote_number.strip(),
            "status": "review_required", "source_system": "EZstimate",
            "pricing_inputs": {k: str(inputs[k]) for k in ("quantity", *FIELDS)},
            "pricing": price, "drawing": drawing, "bom": bom, "purchase_order": po,
            "released_drawing": released_drawing,
            "warnings": warnings,
            "disclaimer": "Estimate proposal; no customer acceptance, purchase order or production release"}
