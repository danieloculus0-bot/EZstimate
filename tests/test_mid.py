"""Synthetic MID workbook adapter validation. No real shop files or supplier prices."""
import io
import os
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from openpyxl import Workbook
from ezstimate.mid_adapter import preview_mid
from ezstimate.server import app


def workbook_bytes():
    wb = Workbook()
    events = wb.active
    events.title = "Forecast Data"
    events.append(["Snapshot Date", "Material Part Number", "Normalized Key", "Need Window",
                   "Window Start", "Window End", "PO Number", "Job Number",
                   "Needed Source Units", "Running Balance Source Units",
                   "Running Balance (Sheets)", "Event Type", "Risk Flag"])
    events.append(["2026-07-23", "SHT-STL-0250X60X120-A36-HR",
                   "SHT-STL-0250X60X120-A36-HR", "07/30 - 08/05",
                   "2026-07-30", "2026-08-05", "", "60194-02", 7200,
                   -7200, -1, "JOB DEMAND", "NEGATIVE BALANCE"])
    queue = wb.create_sheet("MID")
    queue.append(["Action", "Job Due", "Order By", "Material", "Buy Qty",
                  "Supplier", "Controlling Job", "Why Now"])
    queue.append(["BUY NOW - LEAD TIME RISK", "2026-08-05", "2026-07-22",
                  "SHT-STL-0250X60X120-A36-HR", 1, "", "60194-02",
                  "Synthetic material shortage"])
    content = io.BytesIO()
    wb.save(content)
    wb.close()
    return content.getvalue()


class MIDTests(unittest.TestCase):
    def test_reads_job_due_and_negative_balance_as_evidence(self):
        result = preview_mid("mid_demo.xlsx", workbook_bytes())
        self.assertEqual(result["demand_event_count"], 1)
        self.assertEqual(result["buyer_action_count"], 1)
        self.assertEqual(result["buy_now_count"], 1)
        self.assertEqual(result["flagged_event_count"], 1)
        self.assertEqual(result["events"][0]["source_row"], 2)
        self.assertEqual(result["actions"][0]["source_sheet"], "MID")
        self.assertEqual(result["events"][0]["running_balance_sheets"], "-1")
        self.assertIsNone(result["quote_material_price"])
        self.assertFalse(result["persistent_import"])

    def test_rejects_unknown_headers_and_extensions(self):
        with self.assertRaises(ValueError):
            preview_mid("mid_demo.csv", workbook_bytes())
        w=Workbook()
        w.active.append(["Hello","World"])
        f=io.BytesIO()
        w.save(f)
        w.close()
        with self.assertRaisesRegex(ValueError,"No recognized MID"):
            preview_mid("unknown.xlsx", f.getvalue())

    def test_api_preview_does_not_enter_quote_pricing(self):
        with app.test_client() as client:
            result=client.post("/api/pricing/mid-preview",data={
                "file":(io.BytesIO(workbook_bytes()),"mid_demo.xlsx")})
            self.assertEqual(result.status_code,200)
            self.assertEqual(result.json["buy_now_count"],1)
            self.assertIsNone(result.json["quote_material_price"])


if __name__=="__main__":
    unittest.main()
